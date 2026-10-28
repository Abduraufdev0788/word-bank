import csv
import io
import logging
import math
import re
import uuid
from typing import Dict, List, Optional, Tuple
import openpyxl
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.models.vocabulary import ImportJob, Word, WordBatch, WordProgress
from app.schemas.vocabulary import (
    ImportCommitResponse,
    ImportPreviewResponse,
    ImportPreviewRow,
)
from app.services.batch_service import get_or_create_destination_batch, recalculate_batch_status

EN_HEADER_CANDIDATES = {"english", "word", "en", "term", "so'z", "soz", "vocabulary", "eng"}
UZ_HEADER_CANDIDATES = {"uzbek", "translation", "uz", "tarjima", "meaning", "ma'nosi", "manosi", "uzb"}


def detect_columns(headers: List[str]) -> Tuple[Optional[str], Optional[str]]:
    """Detects likely English and Uzbek columns from headers."""
    en_col = None
    uz_col = None
    for h in headers:
        clean = h.strip().lower()
        if not en_col and any(cand in clean for cand in EN_HEADER_CANDIDATES):
            en_col = h
        elif not uz_col and any(cand in clean for cand in UZ_HEADER_CANDIDATES):
            uz_col = h
    return en_col, uz_col


logger = logging.getLogger(__name__)


async def ai_clean_and_extract_words(
    raw_text: str,
) -> List[Tuple[str, str]]:
    """
    Uses Google Gemini AI to analyze, fix, sanitize, and extract English-Uzbek vocabulary
    pairs from ANY arbitrary text, dirty notes, non-standard lists, or mixed documents.
    If English words don't have Uzbek translations in the text, Gemini supplies accurate natural translations.
    """
    if not settings.GEMINI_API_KEY:
        return []

    try:
        from google import genai
        from google.genai import types
        from app.schemas.vocabulary import ExtractedWordsList

        client = genai.Client(api_key=settings.GEMINI_API_KEY)

        sample_text = raw_text[:15000].strip()
        if not sample_text:
            return []

        prompt = (
            "Sen professional leksikograf va lug'at tahrirchisisan.\n"
            "Quyida foydalanuvchi tomonidan yuklangan har xil formatdagi (tartibsiz, belgilari chalkash, qatorlari aralashgan, "
            "raqamlar, defislar yoki hatto faqat inglizcha so'zlar bo'lgan) matn berilgan.\n\n"
            f"Matn:\n{sample_text}\n\n"
            "Vazifang:\n"
            "1. Matndan barcha inglizcha so'z yoki iboralarni (english) va ularning tabiiy to'g'ri o'zbekcha tarjimasini (uzbek) ajratib ol.\n"
            "2. Agar matnda biror inglizcha so'zning o'zbekcha tarjimasi bo'lmasa, unga eng mos o'zbekcha tarjimani o'zing qo'sh.\n"
            "3. Barcha tartib raqamlarini (masalan, 1., 2.), chiziqcha, nuqta, qavslardagi ortiqcha belgilarni tozalab, so'zlarni xatosiz toza holatga keltir.\n"
            "4. Qat'iy ravishda JSON schema bo'yicha 'words' ro'yxatini qaytar."
        )

        def _call_gemini():
            return client.models.generate_content(
                model=settings.GEMINI_MODEL,
                contents=prompt,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json",
                    response_schema=ExtractedWordsList,
                ),
            )

        import asyncio
        response = await asyncio.to_thread(_call_gemini)

        result_words: List[Tuple[str, str]] = []
        parsed_obj = None
        if hasattr(response, "parsed") and response.parsed is not None:
            parsed_obj = response.parsed
        elif getattr(response, "text", None):
            parsed_obj = ExtractedWordsList.model_validate_json(response.text)

        if parsed_obj and parsed_obj.words:
            for item in parsed_obj.words:
                en = item.english.strip().lower()
                uz = item.uzbek.strip()
                if en and uz:
                    result_words.append((en, uz))

        return result_words
    except Exception as e:
        logger.warning(f"AI file sanitization skipped/failed: {e}")
        return []


def parse_txt_content(content_str: str) -> List[Tuple[str, str]]:
    """Parses TXT file lines separated by tab, dash, colon, arrow, equals, or comma."""
    rows = []
    delimiters = [
        "\t",
        " - ",
        " – ",
        " — ",
        " : ",
        " = ",
        " -> ",
        " => ",
        " | ",
        ":",
        "=",
        ",",
    ]
    for line in content_str.splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        found = False
        for delim in delimiters:
            if delim in line:
                parts = line.split(delim, 1)
                en = re.sub(r"^(?:\d+[\.\)]|[-*•])\s*", "", parts[0].strip()).strip().lower()
                uz = re.sub(r"^(?:\d+[\.\)]|[-*•])\s*", "", parts[1].strip()).strip().lower()
                if en and uz:
                    rows.append((en, uz))
                    found = True
                    break
        if not found:
            # Check pattern like "apple (olma)"
            m = re.match(r"^(?:\d+[\.\)]|[-*•])?\s*([a-zA-Z\s\'-]+)\s*\(([^)]+)\)$", line)
            if m:
                en = m.group(1).strip().lower()
                uz = m.group(2).strip().lower()
                if en and uz:
                    rows.append((en, uz))
                    found = True
        if not found and " " in line:
            # Fallback: split on first double space or tab
            parts = re.split(r"\s{2,}", line, maxsplit=1)
            if len(parts) == 2 and parts[0].strip() and parts[1].strip():
                en = re.sub(r"^(?:\d+[\.\)]|[-*•])\s*", "", parts[0].strip()).strip().lower()
                uz = re.sub(r"^(?:\d+[\.\)]|[-*•])\s*", "", parts[1].strip()).strip().lower()
                if en and uz:
                    rows.append((en, uz))
    return rows


def parse_csv_content(
    content_str: str,
    custom_en_col: Optional[str] = None,
    custom_uz_col: Optional[str] = None,
) -> Tuple[List[Tuple[str, str]], Optional[str], Optional[str]]:
    """Parses CSV content with header detection or fallback."""
    # Detect delimiter
    sample = content_str[:2048]
    dialect = csv.excel
    try:
        dialect = csv.Sniffer().sniff(sample)
    except Exception:
        pass

    reader = list(csv.reader(io.StringIO(content_str), dialect))
    if not reader:
        return [], None, None

    headers = [str(col).strip() for col in reader[0]]
    en_col, uz_col = detect_columns(headers)
    if custom_en_col:
        en_col = custom_en_col
    if custom_uz_col:
        uz_col = custom_uz_col

    rows: List[Tuple[str, str]] = []

    if en_col and uz_col and en_col in headers and uz_col in headers:
        en_idx = headers.index(en_col)
        uz_idx = headers.index(uz_col)
        for r in reader[1:]:
            if len(r) > max(en_idx, uz_idx):
                en_val = r[en_idx].strip()
                uz_val = r[uz_idx].strip()
                if en_val and uz_val:
                    rows.append((en_val, uz_val))
    else:
        # Fallback to column 0 and column 1
        start_row = 1 if len(headers) >= 2 and any(c in headers[0].lower() for c in EN_HEADER_CANDIDATES) else 0
        for r in reader[start_row:]:
            if len(r) >= 2:
                en_val = r[0].strip()
                uz_val = r[1].strip()
                if en_val and uz_val:
                    rows.append((en_val, uz_val))
        en_col = en_col or (headers[0] if len(headers) > 0 else "Column 1")
        uz_col = uz_col or (headers[1] if len(headers) > 1 else "Column 2")

    return rows, en_col, uz_col


def parse_xlsx_content(
    file_bytes: bytes,
    custom_en_col: Optional[str] = None,
    custom_uz_col: Optional[str] = None,
) -> Tuple[List[Tuple[str, str]], Optional[str], Optional[str]]:
    """Parses XLSX spreadsheet using openpyxl."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet = wb.active
    if not sheet:
        return [], None, None

    raw_rows = []
    for row in sheet.iter_rows(values_only=True):
        if any(cell is not None for cell in row):
            raw_rows.append([str(c).strip() if c is not None else "" for c in row])

    if not raw_rows:
        return [], None, None

    headers = raw_rows[0]
    en_col, uz_col = detect_columns(headers)
    if custom_en_col:
        en_col = custom_en_col
    if custom_uz_col:
        uz_col = custom_uz_col

    rows: List[Tuple[str, str]] = []
    if en_col and uz_col and en_col in headers and uz_col in headers:
        en_idx = headers.index(en_col)
        uz_idx = headers.index(uz_col)
        for r in raw_rows[1:]:
            if len(r) > max(en_idx, uz_idx):
                en_val = r[en_idx].strip()
                uz_val = r[uz_idx].strip()
                if en_val and uz_val:
                    rows.append((en_val, uz_val))
    else:
        start_row = 1 if len(headers) >= 2 and any(c in headers[0].lower() for c in EN_HEADER_CANDIDATES) else 0
        for r in raw_rows[start_row:]:
            if len(r) >= 2:
                en_val = r[0].strip()
                uz_val = r[1].strip()
                if en_val and uz_val:
                    rows.append((en_val, uz_val))
        en_col = en_col or (headers[0] if len(headers) > 0 else "Column 1")
        uz_col = uz_col or (headers[1] if len(headers) > 1 else "Column 2")

    return rows, en_col, uz_col


async def extract_rows_from_file(
    file_name: str,
    file_bytes: bytes,
    custom_en_col: Optional[str] = None,
    custom_uz_col: Optional[str] = None,
    use_ai: bool = True,
) -> Tuple[List[Tuple[str, str]], str, Optional[str], Optional[str]]:
    """Extracts raw English/Uzbek pairs from TXT, CSV, or XLSX bytes, with AI fallback for arbitrary formats."""
    ext = file_name.split(".")[-1].lower() if "." in file_name else "txt"
    rows: List[Tuple[str, str]] = []

    if ext == "txt":
        content_str = file_bytes.decode("utf-8", errors="replace")
        rows = parse_txt_content(content_str)

        # If standard parser got 0 rows or file has unconventional structure, use AI smart sanitizer
        if (len(rows) == 0 or (use_ai and any(len(line) > 0 and "-" not in line and ":" not in line for line in content_str.splitlines() if line.strip()))) and settings.GEMINI_API_KEY:
            ai_rows = await ai_clean_and_extract_words(content_str)
            if ai_rows:
                return ai_rows, "txt", "English", "Uzbek"

        return rows, "txt", "English", "Uzbek"

    elif ext == "csv":
        content_str = file_bytes.decode("utf-8", errors="replace")
        rows, en_col, uz_col = parse_csv_content(content_str, custom_en_col, custom_uz_col)

        if len(rows) == 0 and use_ai and settings.GEMINI_API_KEY:
            ai_rows = await ai_clean_and_extract_words(content_str)
            if ai_rows:
                return ai_rows, "csv", "English", "Uzbek"

        return rows, "csv", en_col, uz_col

    elif ext in ("xlsx", "xls"):
        rows, en_col, uz_col = parse_xlsx_content(file_bytes, custom_en_col, custom_uz_col)

        if len(rows) == 0 and use_ai and settings.GEMINI_API_KEY:
            # Fallback text representation from xlsx
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
                sheet = wb.active
                lines = []
                if sheet:
                    for r in sheet.iter_rows(values_only=True):
                        lines.append(" - ".join([str(c) for c in r if c is not None]))
                raw_sheet_text = "\n".join(lines)
                ai_rows = await ai_clean_and_extract_words(raw_sheet_text)
                if ai_rows:
                    return ai_rows, "xlsx", "English", "Uzbek"
            except Exception:
                pass

        return rows, "xlsx", en_col, uz_col
    else:
        raise ValueError(f"Qo'llab-quvvatlanmaydigan fayl formati: .{ext}. Faqat TXT, CSV va XLSX mumkin.")


async def generate_import_preview(
    user_id: uuid.UUID,
    file_name: str,
    file_bytes: bytes,
    db: AsyncSession,
    use_ai: bool = True,
) -> ImportPreviewResponse:
    """Generates preview analysis before committing import."""
    rows, file_type, en_col, uz_col = await extract_rows_from_file(file_name, file_bytes, use_ai=use_ai)

    # Get user's existing English words for duplicate checking
    stmt = select(Word.english_word).where(Word.user_id == user_id)
    existing_result = await db.execute(stmt)
    existing_words = {w.lower() for w in existing_result.scalars().all()}

    preview_rows: List[ImportPreviewRow] = []
    duplicate_count = 0
    valid_count = 0

    seen_in_file = set()
    for en, uz in rows:
        en_norm = en.strip().lower()
        uz_norm = uz.strip()
        is_dup = en_norm in existing_words or en_norm in seen_in_file
        seen_in_file.add(en_norm)

        if is_dup:
            duplicate_count += 1
        valid_count += 1

        if len(preview_rows) < 15:
            preview_rows.append(
                ImportPreviewRow(
                    english_word=en_norm,
                    uzbek_translation=uz_norm,
                    is_duplicate=is_dup,
                    is_valid=bool(en_norm and uz_norm),
                )
            )

    net_words = valid_count - duplicate_count
    batches_to_create = math.ceil(net_words / settings.MAX_WORDS_PER_BATCH) if net_words > 0 else 0

    return ImportPreviewResponse(
        file_name=file_name,
        file_type=file_type,
        total_detected_rows=len(rows),
        valid_rows_count=valid_count,
        duplicate_candidates_count=duplicate_count,
        detected_en_column=en_col,
        detected_uz_column=uz_col,
        preview_rows=preview_rows,
        batches_to_create=batches_to_create,
    )


async def execute_commit_import(
    user_id: uuid.UUID,
    file_name: str,
    file_bytes: bytes,
    db: AsyncSession,
    skip_duplicates: bool = True,
    fill_incomplete_batch: bool = True,
    custom_en_col: Optional[str] = None,
    custom_uz_col: Optional[str] = None,
    use_ai: bool = True,
) -> ImportCommitResponse:
    """Commits imported words into 20-word batches transactionally."""
    rows, file_type, _, _ = await extract_rows_from_file(
        file_name, file_bytes, custom_en_col, custom_uz_col, use_ai=use_ai
    )

    # Get user's existing English words
    stmt = select(Word.english_word).where(Word.user_id == user_id)
    existing_result = await db.execute(stmt)
    existing_words = {w.lower() for w in existing_result.scalars().all()}

    # Get or create initial destination batch
    current_batch = await get_or_create_destination_batch(
        user_id=user_id,
        db=db,
        fill_incomplete=fill_incomplete_batch,
        source_file_name=file_name,
    )

    created_batches = set()
    updated_batches = {current_batch.id}
    imported_count = 0
    duplicate_count = 0
    seen_in_session = set()

    for en, uz in rows:
        en_clean = en.strip().lower()
        uz_clean = uz.strip()
        if not en_clean or not uz_clean:
            continue

        if en_clean in existing_words or en_clean in seen_in_session:
            duplicate_count += 1
            if skip_duplicates:
                continue

        # If current batch reached MAX_WORDS_PER_BATCH (20 words), create next batch
        if current_batch.word_count >= settings.MAX_WORDS_PER_BATCH:
            current_batch = await get_or_create_destination_batch(
                user_id=user_id,
                db=db,
                fill_incomplete=False,  # Always create a new batch now
                source_file_name=file_name,
            )
            created_batches.add(current_batch.id)

        # Create Word
        new_word = Word(
            user_id=user_id,
            batch_id=current_batch.id,
            english_word=en_clean,
            uzbek_translation=uz_clean,
            order_in_batch=current_batch.word_count + 1,
        )
        db.add(new_word)
        await db.flush()

        # Initialize WordProgress (SM-2 Spaced Repetition)
        progress = WordProgress(
            user_id=user_id,
            word_id=new_word.id,
            repetition_number=0,
            ease_factor=2.5,
            interval_days=0,
            times_reviewed=0,
            times_correct=0,
            times_incorrect=0,
            is_mastered=False,
            is_weak=False,
        )
        db.add(progress)

        current_batch.word_count += 1
        current_batch.is_complete = current_batch.word_count >= settings.MAX_WORDS_PER_BATCH
        if current_batch.status == "new":
            current_batch.status = "learning"

        seen_in_session.add(en_clean.lower())
        existing_words.add(en_clean.lower())
        imported_count += 1

    # Record ImportJob
    import_job = ImportJob(
        user_id=user_id,
        file_name=file_name,
        file_type=file_type,
        total_rows=len(rows),
        imported_words_count=imported_count,
        duplicate_words_count=duplicate_count,
        failed_rows_count=len(rows) - (imported_count + duplicate_count),
        status="completed",
    )
    db.add(import_job)

    # Recalculate touched batches
    for b_id in updated_batches | created_batches:
        await recalculate_batch_status(b_id, db)

    await db.commit()

    return ImportCommitResponse(
        message=f"{imported_count} ta so'z muvaffaqiyatli import qilindi.",
        total_imported=imported_count,
        total_skipped_duplicates=duplicate_count,
        batches_created=len(created_batches),
        batches_updated=len(updated_batches),
    )
